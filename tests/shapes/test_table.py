# encoding: utf-8

"""Test suite for pptx.table module."""

from __future__ import absolute_import

import pytest

from hamcrest import assert_that, equal_to, is_
from mock import MagicMock, Mock, PropertyMock

from pptx.dml.fill import FillFormat
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml import parse_xml_bytes
from pptx.oxml.ns import nsdecls
from pptx.shapes.table import (
    _Cell, _CellCollection, _Column, _ColumnCollection, _Row, _RowCollection
)
from pptx.util import Inches

from ..oxml.unitdata.shape import test_shapes
from ..oxml.unitdata.table import a_tc, a_tcPr, a_txBody
from ..oxml.unitdata.text import a_bodyPr, a_p, an_r, a_t
from ..unitutil import actual_xml, TestCase


class Describe_Cell(object):

    def it_has_a_fill(self, cell):
        assert isinstance(cell.fill, FillFormat)

    def it_knows_its_margin_settings(self, margin_get_fixture):
        cell, margin_left, margin_right, margin_top, margin_bottom = (
            margin_get_fixture
        )
        assert cell.margin_left == margin_left
        assert cell.margin_right == margin_right
        assert cell.margin_top == margin_top
        assert cell.margin_bottom == margin_bottom

    def it_can_change_its_margin_settings(self, margin_set_fixture):
        cell, marL, marR, marT, marB, tc_with_marX_xml = margin_set_fixture
        cell.margin_left = marL
        cell.margin_right = marR
        cell.margin_top = marT
        cell.margin_bottom = marB
        assert actual_xml(cell._tc) == tc_with_marX_xml

    def it_raises_on_margin_assigned_other_than_int_or_None(
            self, margin_raises_fixture):
        cell, margin_attr_name, val_of_invalid_type = margin_raises_fixture
        with pytest.raises(TypeError):
            setattr(cell, margin_attr_name, val_of_invalid_type)

    def it_can_set_the_text_it_contains(self, text_set_fixture):
        cell, text, tc_with_text_xml = text_set_fixture
        cell.text = text
        assert actual_xml(cell._tc) == tc_with_text_xml

    def it_knows_its_vertical_anchor_setting(self, anchor_get_fixture):
        cell, vertical_anchor = anchor_get_fixture
        assert cell.vertical_anchor == vertical_anchor

    def it_can_change_its_vertical_anchor(self, anchor_set_fixture):
        cell, vertical_anchor, tc_with_anchor_xml = anchor_set_fixture
        cell.vertical_anchor = vertical_anchor
        assert actual_xml(cell._tc) == tc_with_anchor_xml

    # fixture --------------------------------------------------------

    @pytest.fixture(params=[
        (None,  None),
        ('t',   MSO_ANCHOR.TOP),
        ('ctr', MSO_ANCHOR.MIDDLE),
        ('b',   MSO_ANCHOR.BOTTOM),
    ])
    def anchor_get_fixture(self, request):
        anchor, mso_anchor_idx = request.param
        tc_bldr = a_tc().with_nsdecls()
        if anchor is not None:
            tcPr_bldr = a_tcPr().with_anchor(anchor)
            tc_bldr.with_child(tcPr_bldr)
        tc = tc_bldr.element
        cell = _Cell(tc, None)
        return cell, mso_anchor_idx

    @pytest.fixture(params=[
        (None,  None,              None),
        (None,  MSO_ANCHOR.TOP,    't'),
        ('t',   MSO_ANCHOR.MIDDLE, 'ctr'),
        ('ctr', MSO_ANCHOR.BOTTOM, 'b'),
        ('b',   None,              ''),
    ])
    def anchor_set_fixture(self, request):
        def tc_with_anchor_bldr(anchor):
            tc_bldr = a_tc().with_nsdecls()
            if anchor is not None:
                tcPr_bldr = a_tcPr()
                if anchor != '':
                    tcPr_bldr.with_anchor(anchor)
                tc_bldr.with_child(tcPr_bldr)
            return tc_bldr
        before_anchor, mso_anchor_idx, after_anchor = request.param
        tc = tc_with_anchor_bldr(before_anchor).element
        cell = _Cell(tc, None)
        tc_with_anchor_xml = tc_with_anchor_bldr(after_anchor).xml()
        return cell, mso_anchor_idx, tc_with_anchor_xml

    @pytest.fixture
    def cell(self):
        tc = a_tc().with_nsdecls().element
        return _Cell(tc, None)

    @pytest.fixture(params=[
        ((None, None, None, None), (91440, 91440, 45720, 45720)),
        ((1234, 2345, 3456, None), (1234,  2345,  3456,  45720)),
        ((2345, None, 3456, 1234), (2345,  91440, 3456,  1234)),
    ])
    def margin_get_fixture(self, request):
        mar_vals, expected_vals = request.param
        marL, marR, marT, marB = mar_vals
        margin_left, margin_right, margin_top, margin_bottom = expected_vals
        tcPr_bldr = a_tcPr()
        if marL is not None:
            tcPr_bldr.with_marL(marL)
        if marR is not None:
            tcPr_bldr.with_marR(marR)
        if marT is not None:
            tcPr_bldr.with_marT(marT)
        if marB is not None:
            tcPr_bldr.with_marB(marB)
        tc = a_tc().with_nsdecls().with_child(tcPr_bldr).element
        cell = _Cell(tc, None)
        return cell, margin_left, margin_right, margin_top, margin_bottom

    @pytest.fixture(params=[
        ((None, None, None, None), (91440, None,  None, None)),
        ((1234, 2345, 3456, None), (None,  4567,  None, None)),
        ((None, 2345, 3456, None), (None,  None,  6543, None)),
        ((1234, 2345, None, 4567), (None,  None,  None, 7654)),
        ((1234, 2345, 3456, 4567), (None,  None,  None, None)),
        ((None, None, None, None), (None,  None,  None, None)),
    ])
    def margin_set_fixture(self, request):

        def tc_bldr_with(marX):
            marL, marR, marT, marB = marX
            tcPr_bldr = a_tcPr()
            if marL is not None:
                tcPr_bldr.with_marL(marL)
            if marR is not None:
                tcPr_bldr.with_marR(marR)
            if marT is not None:
                tcPr_bldr.with_marT(marT)
            if marB is not None:
                tcPr_bldr.with_marB(marB)
            return a_tc().with_nsdecls().with_child(tcPr_bldr)

        before_vals, after_vals = request.param
        tc = tc_bldr_with(before_vals).element
        cell = _Cell(tc, None)

        marL, marR, marT, marB = after_vals

        tc_with_marX_xml = tc_bldr_with(after_vals).xml()

        return cell, marL, marR, marT, marB, tc_with_marX_xml

    @pytest.fixture(params=[
        'margin_left', 'margin_right', 'margin_top', 'margin_bottom'
    ])
    def margin_raises_fixture(self, request, cell):
        margin_attr_name = request.param
        val_of_invalid_type = 'foobar'
        return cell, margin_attr_name, val_of_invalid_type

    @pytest.fixture
    def text_set_fixture(self, cell):
        text = 'foobar'
        tc_with_text_xml = (
            a_tc().with_nsdecls().with_child(
                a_txBody().with_child(
                    a_bodyPr()).with_child(
                    a_p().with_child(
                        an_r().with_child(
                            a_t().with_text(text)))))
            .xml()
        )
        return cell, text, tc_with_text_xml


class Test_CellCollection(TestCase):
    """Test _CellCollection"""
    def setUp(self):
        tr_xml = (
            '<a:tr %s h="370840"><a:tc><a:txBody><a:p/></a:txBody></a:tc><a:t'
            'c><a:txBody><a:p/></a:txBody></a:tc></a:tr>' % nsdecls('a')
        )
        test_tr_elm = parse_xml_bytes(tr_xml)
        self.cells = _CellCollection(test_tr_elm, None)

    def test_is_indexable(self):
        """_CellCollection indexable (e.g. no TypeError on 'cells[0]')"""
        # verify -----------------------
        try:
            self.cells[0]
        except TypeError:
            msg = "'_CellCollection' object does not support indexing"
            self.fail(msg)
        except IndexError:
            pass

    def test_is_iterable(self):
        """_CellCollection is iterable (e.g. ``for cell in cells:``)"""
        # verify -----------------------
        count = 0
        try:
            for cell in self.cells:
                count += 1
        except TypeError:
            msg = "_CellCollection object is not iterable"
            self.fail(msg)
        assert_that(count, is_(equal_to(2)))

    def test_raises_on_idx_out_of_range(self):
        """_CellCollection raises on index out of range"""
        with self.assertRaises(IndexError):
            self.cells[9]

    def test_cell_count_correct(self):
        """len(_CellCollection) returns correct cell count"""
        # verify -----------------------
        assert_that(len(self.cells), is_(equal_to(2)))


class Test_Column(TestCase):
    """Test _Column"""
    def setUp(self):
        gridCol_xml = '<a:gridCol %s w="3048000"/>' % nsdecls('a')
        test_gridCol_elm = parse_xml_bytes(gridCol_xml)
        self.column = _Column(test_gridCol_elm, Mock(name='table'))

    def test_width_from_xml_correct(self):
        """_Column.width returns correct value from gridCol XML element"""
        # verify -----------------------
        assert_that(self.column.width, is_(equal_to(3048000)))

    def test_width_round_trips_intact(self):
        """_Column.width round-trips intact"""
        # setup ------------------------
        self.column.width = 999
        # verify -----------------------
        assert_that(self.column.width, is_(equal_to(999)))

    def test_set_width_raises_on_bad_value(self):
        """_Column.width raises on attempt to assign invalid value"""
        test_cases = ('abc', '1', -1)
        for value in test_cases:
            with self.assertRaises(ValueError):
                self.column.width = value


class Test_ColumnCollection(TestCase):
    """Test _ColumnCollection"""
    def setUp(self):
        tbl_xml = (
            '<a:tbl %s><a:tblGrid><a:gridCol w="3048000"/><a:gridCol w="30480'
            '00"/></a:tblGrid></a:tbl>' % nsdecls('a')
        )
        test_tbl_elm = parse_xml_bytes(tbl_xml)
        self.columns = _ColumnCollection(test_tbl_elm, Mock(name='table'))

    def test_is_indexable(self):
        """_ColumnCollection indexable (e.g. no TypeError on 'columns[0]')"""
        # verify -----------------------
        try:
            self.columns[0]
        except TypeError:
            msg = "'_ColumnCollection' object does not support indexing"
            self.fail(msg)
        except IndexError:
            pass

    def test_is_iterable(self):
        """_ColumnCollection is iterable (e.g. ``for col in columns:``)"""
        # verify -----------------------
        count = 0
        try:
            for col in self.columns:
                count += 1
        except TypeError:
            msg = "_ColumnCollection object is not iterable"
            self.fail(msg)
        assert_that(count, is_(equal_to(2)))

    def test_raises_on_idx_out_of_range(self):
        """_ColumnCollection raises on index out of range"""
        with self.assertRaises(IndexError):
            self.columns[9]

    def test_column_count_correct(self):
        """len(_ColumnCollection) returns correct column count"""
        # verify -----------------------
        assert_that(len(self.columns), is_(equal_to(2)))


class Test_Row(TestCase):
    """Test _Row"""
    def setUp(self):
        tr_xml = (
            '<a:tr %s h="370840"><a:tc><a:txBody><a:p/></a:txBody></a:tc><a:t'
            'c><a:txBody><a:p/></a:txBody></a:tc></a:tr>' % nsdecls('a')
        )
        test_tr_elm = parse_xml_bytes(tr_xml)
        self.row = _Row(test_tr_elm, Mock(name='table'))

    def test_height_from_xml_correct(self):
        """_Row.height returns correct value from tr XML element"""
        # verify -----------------------
        assert_that(self.row.height, is_(equal_to(370840)))

    def test_height_round_trips_intact(self):
        """_Row.height round-trips intact"""
        # setup ------------------------
        self.row.height = 999
        # verify -----------------------
        assert_that(self.row.height, is_(equal_to(999)))

    def test_set_height_raises_on_bad_value(self):
        """_Row.height raises on attempt to assign invalid value"""
        test_cases = ('abc', '1', -1)
        for value in test_cases:
            with self.assertRaises(ValueError):
                self.row.height = value


class Test_RowCollection(TestCase):
    """Test _RowCollection"""
    def setUp(self):
        tbl_xml = (
            '<a:tbl %s><a:tr h="370840"><a:tc><a:txBody><a:p/></a:txBody></a:'
            'tc><a:tc><a:txBody><a:p/></a:txBody></a:tc></a:tr><a:tr h="37084'
            '0"><a:tc><a:txBody><a:p/></a:txBody></a:tc><a:tc><a:txBody><a:p/'
            '></a:txBody></a:tc></a:tr></a:tbl>' % nsdecls('a')
        )
        test_tbl_elm = parse_xml_bytes(tbl_xml)
        self.rows = _RowCollection(test_tbl_elm, Mock(name='table'))

    def test_is_indexable(self):
        """_RowCollection indexable (e.g. no TypeError on 'rows[0]')"""
        # verify -----------------------
        try:
            self.rows[0]
        except TypeError:
            msg = "'_RowCollection' object does not support indexing"
            self.fail(msg)
        except IndexError:
            pass

    def test_is_iterable(self):
        """_RowCollection is iterable (e.g. ``for row in rows:``)"""
        # verify -----------------------
        count = 0
        try:
            for row in self.rows:
                count += 1
        except TypeError:
            msg = "_RowCollection object is not iterable"
            self.fail(msg)
        assert_that(count, is_(equal_to(2)))

    def test_raises_on_idx_out_of_range(self):
        """_RowCollection raises on index out of range"""
        with self.assertRaises(IndexError):
            self.rows[9]

    def test_row_count_correct(self):
        """len(_RowCollection) returns correct row count"""
        # verify -----------------------
        assert_that(len(self.rows), is_(equal_to(2)))


class TestTable(TestCase):
    """Test Table"""
    def test_initial_height_divided_evenly_between_rows(self):
        """Table creation height divided evenly between rows"""
        # constant values -------------
        rows = cols = 3
        left = top = Inches(1.0)
        width = Inches(2.0)
        height = 1000
        shapes = test_shapes.empty_shape_collection
        # exercise ---------------------
        table = shapes.add_table(rows, cols, left, top, width, height)
        # verify -----------------------
        assert_that(table.rows[0].height, is_(equal_to(333)))
        assert_that(table.rows[1].height, is_(equal_to(333)))
        assert_that(table.rows[2].height, is_(equal_to(334)))

    def test_initial_width_divided_evenly_between_columns(self):
        """Table creation width divided evenly between columns"""
        # constant values -------------
        rows = cols = 3
        left = top = Inches(1.0)
        width = 1000
        height = Inches(2.0)
        shapes = test_shapes.empty_shape_collection
        # exercise ---------------------
        table = shapes.add_table(rows, cols, left, top, width, height)
        # verify -----------------------
        assert_that(table.columns[0].width, is_(equal_to(333)))
        assert_that(table.columns[1].width, is_(equal_to(333)))
        assert_that(table.columns[2].width, is_(equal_to(334)))

    def test_height_sum_of_row_heights(self):
        """Table.height is sum of row heights"""
        # constant values -------------
        rows = cols = 2
        left = top = width = height = Inches(2.0)
        # setup ------------------------
        shapes = test_shapes.empty_shape_collection
        tbl = shapes.add_table(rows, cols, left, top, width, height)
        tbl.rows[0].height = 100
        tbl.rows[1].height = 200
        # verify -----------------------
        sum_of_row_heights = 300
        assert_that(tbl.height, is_(equal_to(sum_of_row_heights)))

    def test_width_sum_of_col_widths(self):
        """Table.width is sum of column widths"""
        # constant values -------------
        rows = cols = 2
        left = top = width = height = Inches(2.0)
        # setup ------------------------
        shapes = test_shapes.empty_shape_collection
        tbl = shapes.add_table(rows, cols, left, top, width, height)
        tbl.columns[0].width = 100
        tbl.columns[1].width = 200
        # verify -----------------------
        sum_of_col_widths = tbl.columns[0].width + tbl.columns[1].width
        assert_that(tbl.width, is_(equal_to(sum_of_col_widths)))


class TestTableBooleanProperties(TestCase):
    """Test Table"""
    def setUp(self):
        """Test fixture for Table boolean properties"""
        shapes = test_shapes.empty_shape_collection
        self.table = shapes.add_table(2, 2, 1000, 1000, 1000, 1000)
        self.assignment_cases = (
            (True,  True),
            (False, False),
            (0,     False),
            (1,     True),
            ('',    False),
            ('foo', True)
        )

    def mockery(self, property_name, property_return_value=None):
        """
        Return property of *property_name* on self.table with return value of
        *property_return_value*.
        """
        # mock <a:tbl> element of Table so we can mock its properties
        tbl = MagicMock()
        self.table._tbl_elm = tbl
        # create a suitable mock for the property
        property_ = PropertyMock()
        if property_return_value:
            property_.return_value = property_return_value
        # and attach it the the <a:tbl> element object (class actually)
        setattr(type(tbl), property_name, property_)
        return property_

    def test_first_col_property_value(self):
        """Table.first_col property value is calculated correctly"""
        # mockery ----------------------
        firstCol_val = True
        firstCol = self.mockery('firstCol', firstCol_val)
        # exercise ---------------------
        retval = self.table.first_col
        # verify -----------------------
        firstCol.assert_called_once_with()
        assert_that(retval, is_(equal_to(firstCol_val)))

    def test_first_row_property_value(self):
        """Table.first_row property value is calculated correctly"""
        # mockery ----------------------
        firstRow_val = True
        firstRow = self.mockery('firstRow', firstRow_val)
        # exercise ---------------------
        retval = self.table.first_row
        # verify -----------------------
        firstRow.assert_called_once_with()
        assert_that(retval, is_(equal_to(firstRow_val)))

    def test_horz_banding_property_value(self):
        """Table.horz_banding property value is calculated correctly"""
        # mockery ----------------------
        bandRow_val = True
        bandRow = self.mockery('bandRow', bandRow_val)
        # exercise ---------------------
        retval = self.table.horz_banding
        # verify -----------------------
        bandRow.assert_called_once_with()
        assert_that(retval, is_(equal_to(bandRow_val)))

    def test_last_col_property_value(self):
        """Table.last_col property value is calculated correctly"""
        # mockery ----------------------
        lastCol_val = True
        lastCol = self.mockery('lastCol', lastCol_val)
        # exercise ---------------------
        retval = self.table.last_col
        # verify -----------------------
        lastCol.assert_called_once_with()
        assert_that(retval, is_(equal_to(lastCol_val)))

    def test_last_row_property_value(self):
        """Table.last_row property value is calculated correctly"""
        # mockery ----------------------
        lastRow_val = True
        lastRow = self.mockery('lastRow', lastRow_val)
        # exercise ---------------------
        retval = self.table.last_row
        # verify -----------------------
        lastRow.assert_called_once_with()
        assert_that(retval, is_(equal_to(lastRow_val)))

    def test_vert_banding_property_value(self):
        """Table.vert_banding property value is calculated correctly"""
        # mockery ----------------------
        bandCol_val = True
        bandCol = self.mockery('bandCol', bandCol_val)
        # exercise ---------------------
        retval = self.table.vert_banding
        # verify -----------------------
        bandCol.assert_called_once_with()
        assert_that(retval, is_(equal_to(bandCol_val)))

    def test_first_col_assignment(self):
        """Assignment to Table.first_col sets attribute value"""
        # mockery ----------------------
        firstCol = self.mockery('firstCol')
        # verify -----------------------
        for assigned_value, called_with_value in self.assignment_cases:
            self.table.first_col = assigned_value
            firstCol.assert_called_once_with(called_with_value)
            firstCol.reset_mock()

    def test_first_row_assignment(self):
        """Assignment to Table.first_row sets attribute value"""
        # mockery ----------------------
        firstRow = self.mockery('firstRow')
        # verify -----------------------
        for assigned_value, called_with_value in self.assignment_cases:
            self.table.first_row = assigned_value
            firstRow.assert_called_once_with(called_with_value)
            firstRow.reset_mock()

    def test_horz_banding_assignment(self):
        """Assignment to Table.horz_banding sets attribute value"""
        # mockery ----------------------
        bandRow = self.mockery('bandRow')
        # verify -----------------------
        for assigned_value, called_with_value in self.assignment_cases:
            self.table.horz_banding = assigned_value
            bandRow.assert_called_once_with(called_with_value)
            bandRow.reset_mock()

    def test_last_col_assignment(self):
        """Assignment to Table.last_col sets attribute value"""
        # mockery ----------------------
        lastCol = self.mockery('lastCol')
        # verify -----------------------
        for assigned_value, called_with_value in self.assignment_cases:
            self.table.last_col = assigned_value
            lastCol.assert_called_once_with(called_with_value)
            lastCol.reset_mock()

    def test_last_row_assignment(self):
        """Assignment to Table.last_row sets attribute value"""
        # mockery ----------------------
        lastRow = self.mockery('lastRow')
        # verify -----------------------
        for assigned_value, called_with_value in self.assignment_cases:
            self.table.last_row = assigned_value
            lastRow.assert_called_once_with(called_with_value)
            lastRow.reset_mock()

    def test_vert_banding_assignment(self):
        """Assignment to Table.vert_banding sets attribute value"""
        # mockery ----------------------
        bandCol = self.mockery('bandCol')
        # verify -----------------------
        for assigned_value, called_with_value in self.assignment_cases:
            self.table.vert_banding = assigned_value
            bandCol.assert_called_once_with(called_with_value)
            bandCol.reset_mock()
